"""Build the compact Jenkar-style Carbone Booking confirmation Word template."""
from pathlib import Path
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt, RGBColor

DESTINATION = Path(__file__).with_name('Booking_Information_Carbone_Template.docx')
BLUE = RGBColor(38, 72, 96)
INK = RGBColor(35, 46, 54)
WIDTH = 180

doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Mm(210), Mm(297)
sec.top_margin, sec.bottom_margin = Mm(12), Mm(13)
sec.left_margin = sec.right_margin = Mm(15)
sec.footer_distance = Mm(6)
normal = doc.styles['Normal']
normal.font.name, normal.font.size, normal.font.color.rgb = 'Arial', Pt(8), INK
normal.paragraph_format.space_after = Pt(0)
normal.paragraph_format.line_spacing = 1.03


def border(cell):
    props = cell._tc.get_or_add_tcPr()
    edges = props.find(qn('w:tcBorders'))
    if edges is None:
        edges = OxmlElement('w:tcBorders')
        props.append(edges)
    for side in ('top', 'left', 'bottom', 'right'):
        line = OxmlElement('w:' + side)
        line.set(qn('w:val'), 'single')
        line.set(qn('w:sz'), '4')
        line.set(qn('w:color'), '7E8D97')
        edges.append(line)
    mar = OxmlElement('w:tcMar')
    for side, value in (('top', 55), ('bottom', 55), ('start', 85), ('end', 85)):
        item = OxmlElement('w:' + side)
        item.set(qn('w:w'), str(value))
        item.set(qn('w:type'), 'dxa')
        mar.append(item)
    props.append(mar)


def grid(widths, rows):
    assert sum(widths) == WIDTH
    t = doc.add_table(rows=rows, cols=len(widths))
    t.autofit = False
    for column, width in zip(t.columns, widths):
        column.width = Mm(width)
    for row in t.rows:
        for cell, width in zip(row.cells, widths):
            cell.width = Mm(width)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            border(cell)
    return t


def cell_content(cell, label, value):
    cell.text = ''
    p = cell.paragraphs[0]
    r = p.add_run(label.upper())
    r.bold = True
    r.font.size = Pt(6.1)
    r.font.color.rgb = BLUE
    r = p.add_run('\n' + value)
    r.font.size = Pt(7.1)


def alias(mapping):
    p = doc.add_paragraph()
    p.paragraph_format.line_spacing = 0.1
    for name, expression in mapping.items():
        p.add_run('{#' + name + '=' + expression + '}').font.size = Pt(1)


def gate(tag):
    p = doc.add_paragraph(tag)
    p.paragraph_format.line_spacing = 0.1
    for run in p.runs:
        run.font.size = Pt(1)


def section(text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run(text.upper())
    r.bold = True
    r.font.size = Pt(7)
    r.font.color.rgb = BLUE


def item_table(headers, values, marker, widths):
    t = grid(widths, 1)
    header_properties = t.rows[0]._tr.get_or_add_trPr()
    repeat = OxmlElement('w:tblHeader')
    repeat.set(qn('w:val'), 'true')
    header_properties.append(repeat)
    for cell, title in zip(t.rows[0].cells, headers):
        cell.text = title.upper()
        for run in cell.paragraphs[0].runs:
            run.bold = True
            run.font.size = Pt(6.1)
            run.font.color.rgb = BLUE
    cells = t.add_row().cells
    row_properties = t.rows[1]._tr.get_or_add_trPr()
    no_split = OxmlElement('w:cantSplit')
    row_properties.append(no_split)
    for cell, width, value in zip(cells, widths, values):
        cell.width = Mm(width)
        border(cell)
        cell.text = value
        for run in cell.paragraphs[0].runs:
            run.font.size = Pt(7)
    end = t.add_row().cells
    for cell, width in zip(end, widths):
        cell.width = Mm(width)
        cell.text = ''
    end[0].text = marker


alias({
    'ref': 'd.bookingConfirmation.bookingReference',
    'direction': 'd.bookingConfirmation.direction',
    'mode': 'd.bookingConfirmation.mode',
    'shipmentType': 'd.bookingConfirmation.shipmentType',
    'incoterm': 'd.bookingConfirmation.incoterm',
    'customer': 'd.customer.name',
    'customerRef': 'd.bookingConfirmation.customerReference',
    'shipper': 'd.bookingConfirmation.shipper.name',
    'shipperAddress': 'd.bookingConfirmation.shipper.address',
    'consignee': 'd.bookingConfirmation.consignee.name',
    'consigneeAddress': 'd.bookingConfirmation.consignee.address',
    'collectionAddress': 'd.bookingConfirmation.collection.address',
    'collectionDate': 'd.bookingConfirmation.collection.plannedAtLabel',
    'collectionRemarks': 'd.bookingConfirmation.collection.remarks',
    'deliveryAddress': 'd.bookingConfirmation.delivery.address',
    'deliveryDate': 'd.bookingConfirmation.delivery.plannedAtLabel',
    'deliveryRemarks': 'd.bookingConfirmation.delivery.remarks',
    'preparedBy': 'd.bookingConfirmation.preparedBy',
    'instructions': 'd.bookingConfirmation.specialInstructions',
})

head = doc.add_table(rows=1, cols=2)
head.autofit = False
head.columns[0].width, head.columns[1].width = Mm(115), Mm(65)
left, right = head.rows[0].cells
left.width, right.width = Mm(115), Mm(65)
left.text = right.text = ''
r = left.paragraphs[0].add_run('JENKAR')
r.bold, r.font.name, r.font.size, r.font.color.rgb = True, 'Arial', Pt(23), BLUE
right.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT
r = right.paragraphs[0].add_run('Jenkar Shipping Limited')
r.bold, r.font.size = True, Pt(8)

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_before = Pt(3)
p.paragraph_format.space_after = Pt(4)
r = p.add_run('Booking Confirmation  ·  Our Ref: {$ref}')
r.bold, r.font.size = True, Pt(11)

gate("{d.bookingConfirmation.provisional:ifEQ(true):showBegin}")
p = doc.add_paragraph('PROVISIONAL BOOKING — NOT YET CONFIRMED')
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
for r in p.runs:
    r.bold, r.font.size = True, Pt(7)
gate('{d.bookingConfirmation.provisional:showEnd}')

p = doc.add_paragraph("{d.bookingConfirmation.provisional:ifEQ(true):show('Please review the provisional transport arrangements below.'):elseShow('This confirms the transport arrangements shown below.')} Please check the details and tell us promptly if anything needs correcting.")
p.paragraph_format.space_after = Pt(7)

parties = grid([90, 90], 2)
cell_content(parties.cell(0, 0), 'Shipper', '{$shipper}\n{$shipperAddress}')
cell_content(parties.cell(0, 1), 'Pick up from', '{$collectionAddress}\n{$collectionDate}')
cell_content(parties.cell(1, 0), 'Consignee', '{$consignee}\n{$consigneeAddress}')
cell_content(parties.cell(1, 1), 'Customer', '{$customer}\nReference: {$customerRef}')

section('Booking information')
info = grid([45, 45, 45, 45], 2)
for cell, label, value in zip(info.rows[0].cells,
        ['Booking reference', 'Direction', 'Mode', 'Shipment type'],
        ['{$ref}', '{$direction}', '{$mode}', '{$shipmentType}']):
    cell_content(cell, label, value)
for cell, label, value in zip(info.rows[1].cells,
        ['Customer reference', 'Incoterms', 'Collection', 'Delivery'],
        ['{$customerRef}', '{$incoterm}', '{$collectionDate}', '{$deliveryDate}']):
    cell_content(cell, label, value)

gate('{d.bookingConfirmation.scope.mainTransport:ifEQ(true):showBegin}')
alias({
    'transport': 'd.bookingConfirmation.mainTransport',
    'from': 'd.bookingConfirmation.mainTransport[i].origin',
    'to': 'd.bookingConfirmation.mainTransport[i].destination',
    # Keep the UTC clock time from the server snapshot. Local date formatting
    # shifts it according to the renderer host's timezone.
    'departureDate': 'd.bookingConfirmation.mainTransport[i].plannedDepartureAt:substr(0,10)',
    'departureTime': 'd.bookingConfirmation.mainTransport[i].plannedDepartureAt:substr(11,16)',
    'arrivalDate': 'd.bookingConfirmation.mainTransport[i].plannedArrivalAt:substr(0,10)',
    'arrivalTime': 'd.bookingConfirmation.mainTransport[i].plannedArrivalAt:substr(11,16)',
    'routeDetails': 'd.bookingConfirmation.mainTransport[i].details',
})
section('Main transport')
item_table(['Mode', 'From / to', 'Planned UTC', 'Service and references'],
    ['{$transport[i].mode}', '{$from}\n→ {$to}',
     'Dep {$departureDate} {$departureTime}\nArr {$arrivalDate} {$arrivalTime}', '{$routeDetails}'],
    '{$transport[i+1]}', [18, 47, 45, 70])
gate('{d.bookingConfirmation.scope.mainTransport:showEnd}')

gate('{d.bookingConfirmation.hasEquipment:ifEQ(true):showBegin}')
alias({'equipment': 'd.bookingConfirmation.equipment'})
section('Transport equipment')
item_table(['Kind', 'Number', 'Type'],
    ['{$equipment[i].kind}', '{$equipment[i].number}', '{$equipment[i].type}'],
    '{$equipment[i+1]}', [35, 95, 50])
gate('{d.bookingConfirmation.hasEquipment:showEnd}')

alias({
    'goods': 'd.bookingConfirmation.cargo',
    'marks': 'd.bookingConfirmation.cargo[i].marksAndNumbers',
    'packages': 'd.bookingConfirmation.cargo[i].packages',
    'packageType': 'd.bookingConfirmation.cargo[i].packageType',
    'weight': 'd.bookingConfirmation.cargo[i].grossWeightKg',
})
section('Shipment information')
item_table(['Description', 'Marks and numbers', 'Packages', 'Type', 'Weight kg'],
    ['{$goods[i].description}', '{$marks}', '{$packages}', '{$packageType}', '{$weight}'],
    '{$goods[i+1]}', [57, 42, 19, 30, 32])
p = doc.add_paragraph()
p.paragraph_format.space_before = Pt(3)
p.add_run('Special instructions  ').bold = True
p.add_run('{$instructions}')
p = doc.add_paragraph('Prepared by {$preparedBy}')
p.paragraph_format.space_before = Pt(2)

section('Customer selling charges')
gate("{d.bookingConfirmation.priceStatus:ifNE('confirmed'):showBegin}")
doc.add_paragraph('Price to be confirmed')
gate('{d.bookingConfirmation.priceStatus:showEnd}')
gate("{d.bookingConfirmation.priceStatus:ifEQ('confirmed'):showBegin}")
alias({
    'charges': 'd.bookingConfirmation.chargeLines',
    'chargeCurrency': 'd.bookingConfirmation.chargeLines[i].currency',
    'chargeAmount': 'd.bookingConfirmation.chargeLines[i].sellAmount:formatN(2)',
    'totals': 'd.bookingConfirmation.chargeTotals',
    'totalAmount': 'd.bookingConfirmation.chargeTotals[i].amount:formatN(2)',
})
item_table(['Service', 'Currency', 'Selling price'],
    ['{$charges[i].description}', '{$chargeCurrency}', '{$chargeAmount}'],
    '{$charges[i+1]}', [116, 28, 36])
item_table(['Total currency', 'Customer total'],
    ['{$totals[i].currency}', '{$totalAmount}'],
    '{$totals[i+1]}', [90, 90])
gate('{d.bookingConfirmation.priceStatus:showEnd}')

footer = sec.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = footer.add_run('Jenkar Shipping Limited  ·  Booking {$ref}')
r.font.size, r.font.color.rgb = Pt(6.7), BLUE

doc.save(DESTINATION)
print(DESTINATION)
