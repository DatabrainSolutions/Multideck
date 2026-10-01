#!/usr/bin/env python3
"""Separate Draft-only revisions. Retain every original package part except body XML.

No completed reference is used: these inputs are the previously privacy-reviewed
clean sources. Form geometry/artwork and terms remain unchanged. Unsafe populated
slots become blank or point to an unbounded inline schedule (not textbox loops).
"""
from copy import deepcopy
from io import BytesIO
from pathlib import Path
import hashlib
import json
import zipfile

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt
from lxml import etree

ROOT = Path(__file__).resolve().parents[3]
OUT = Path(__file__).resolve().parent
W = 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'
NS = {'w': W, 'v': 'urn:schemas-microsoft-com:vml'}


def schedule():
    document = Document()
    section = document.sections[0]
    section.page_width, section.page_height = Mm(210), Mm(297)
    section.top_margin = section.bottom_margin = Mm(18)
    section.left_margin = section.right_margin = Mm(16)
    normal = document.styles['Normal']
    normal.font.name, normal.font.size = 'Arial', Pt(9)
    normal.paragraph_format.space_after = Pt(6)
    document.add_paragraph('DRAFT — TRANSPORT PARTICULARS FOR REVIEW').runs[0].bold = True
    document.add_paragraph('Booking: {d.transport.reference}')
    document.add_paragraph('This schedule forms part of the Draft. Blank values are not recorded. '
                           'No security clearance, verified mass, signature or issuing authority is implied.')

    def heading(text):
        paragraph = document.add_paragraph(text)
        paragraph.paragraph_format.keep_with_next = True
        paragraph.runs[0].bold = True

    def repeat_table(title, path, columns):
        heading(title)
        table = document.add_table(rows=1, cols=len(columns))
        table.autofit = False
        available = 178
        for column, (_, _, proportion) in zip(table.columns, columns):
            column.width = Mm(available * proportion)
        header = table.rows[0]
        header._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'))
        for cell, (label, _, proportion) in zip(header.cells, columns):
            cell.width = Mm(available * proportion)
            cell.text = label
            cell.paragraphs[0].runs[0].bold = True
        # One normal-flow repeated row + its next-index boundary. No fixed
        # heights, clipping, ellipsis or floating drawings on any schedule row.
        for index in ('i', 'i+1'):
            row = table.add_row()
            for cell, (_, field, proportion) in zip(row.cells, columns):
                cell.width = Mm(available * proportion)
                cell.text = '{d.' + path + '[' + index + '].' + field + ':convCRLF}'
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    paragraph.paragraph_format.space_after = Pt(4)
                    paragraph.paragraph_format.space_before = Pt(4)
                    for run in paragraph.runs:
                        run.font.name, run.font.size = 'Arial', Pt(8)
                borders = OxmlElement('w:tcBorders')
                edge = OxmlElement('w:bottom')
                edge.set(qn('w:val'), 'single'); edge.set(qn('w:sz'), '3'); edge.set(qn('w:color'), 'BBBBBB')
                borders.append(edge); cell._tc.get_or_add_tcPr().append(borders)

    repeat_table('Parties — full saved names and addresses', 'transport.parties', [
        ('Role', 'role', .18), ('Name', 'name', .32), ('Full address', 'fullAddress', .5)])
    repeat_table('Cargo — all lines', 'transport.cargo', [
        ('Line', 'lineNumber', .06), ('Goods', 'description', .29), ('Packages', 'packageQuantity', .12),
        ('Pack type', 'packageType', .11), ('Gross kg', 'grossWeight', .12), ('Chargeable kg', 'chargeableWeight', .15), ('CBM', 'volume', .15)])
    repeat_table('Cargo marks, commodity and handling', 'transport.cargo', [
        ('Line', 'lineNumber', .06), ('Marks', 'marksAndNumbers', .27), ('Commodity', 'commodity', .27), ('HS code', 'hsCode', .13), ('Handling', 'handling', .27)])
    repeat_table('Equipment — all containers or ULDs', 'transport.equipment', [
        ('Number', 'number', .24), ('Type', 'type', .12), ('Seal', 'seal', .18), ('Recorded gross kg', 'grossWeight', .16),
        ('Recorded VGM kg', 'verifiedGrossMass', .16), ('CBM', 'volume', .14)])
    repeat_table('Cargo splits — recorded quantities only', 'transport.allocations', [
        ('Cargo line', 'cargoLine', .12), ('Equipment', 'equipmentNumber', .25), ('Type', 'equipmentType', .15),
        ('Packages', 'packages', .16), ('Allocated gross kg', 'grossWeight', .16), ('Allocated CBM', 'volume', .16)])
    document.add_paragraph('Package quantities do not calculate container weight or VGM. '
                           'Unnumbered equipment remains listed. Unrecorded splits are not estimated.')
    repeat_table('Information still to review', 'transport.gaps', [('Not recorded / review required', 'label', 1)])
    for paragraph in document.paragraphs:
        paragraph.paragraph_format.line_spacing = 1.1
        paragraph.paragraph_format.space_after = Pt(4)
        for run in paragraph.runs:
            run.font.name, run.font.size = 'Arial', Pt(9)
    return document.element.body


def set_text(paragraph, value):
    nodes = paragraph.xpath('.//w:t', namespaces=NS)
    if nodes:
        nodes[0].text = value
        for node in nodes[1:]:
            node.text = ''


def patch_air(root):
    for shape in root.xpath('//v:shape[v:textbox]', namespaces=NS):
        name = shape.get('id', '')
        field = name.split('-', 2)[-1]
        paragraphs = shape.xpath('.//w:txbxContent/w:p', namespaces=NS)
        replacements = {
            'agent': ['{d.job.legalEntityName:ellipsis(46)}', ''],
            'handling': ['DRAFT - SEE ATTACHED CARGO SCHEDULE', ''], 'sci': [''],
            'cargo-pieces': ['{d.transport.totals.packages}'], 'total-pieces': ['{d.transport.totals.packages}'],
            'cargo-weight': ['{d.transport.totals.grossWeight}'], 'total-weight': ['{d.transport.totals.grossWeight}'],
            'cargo-chargeable': ['{d.transport.totals.chargeableWeight}'], 'total-chargeable': ['{d.transport.totals.chargeableWeight}'],
            'cargo-nature': ['SEE ATTACHED', 'CARGO SCHEDULE', 'FULL PARTICULARS'],
            'cargo-item': [''], 'charges-description': [''],
            'total-volume': ['{d.transport.totals.volume} M3'], 'signature': ['', ''], 'issue-detail': [''], 'issuer-signature': [''],
        }
        if field in replacements:
            for paragraph, value in zip(paragraphs, replacements[field]):
                set_text(paragraph, value)
        elif field == 'copy-label':
            label = ''.join(paragraphs[0].xpath('.//w:t/text()', namespaces=NS))
            set_text(paragraphs[0], 'DRAFT - ' + label.replace('ORIGINAL ', 'FORM ').replace('COPY ', 'FORM '))
        elif field in ('origin-code', 'route-to'):
            endpoint = 'origin' if field == 'origin-code' else 'destination'
            set_text(paragraphs[0], '{d.routing[isMainCarriage=true].' + endpoint + '.iataCode}')
    # The MNG retained form uses the same mawb-... field identifiers.
    for paragraph in root.xpath('//w:p', namespaces=NS):
        text = ''.join(paragraph.xpath('./w:r/w:t/text()', namespaces=NS))
        if text.startswith('SHP REF:'):
            set_text(paragraph, 'BOOKING REF: {d.transport.reference:ellipsis(28)}')
        elif text.startswith(('ORIGINAL ', 'COPY ')):
            set_text(paragraph, 'DRAFT - ' + text.replace('ORIGINAL ', 'FORM ').replace('COPY ', 'FORM '))


def patch_fiata(root):
    # This Draft layout is Multideck-owned. Retain its inspected default logo
    # rather than dropping it when a Booking has no authorised brand mapping.
    # The image remains directly replaceable in the editable Word layout.
    for node in root.iter():
        for attribute in ('descr', 'title', 'alt'):
            if 'company.logoDataUri' in node.get(attribute, ''):
                node.set(attribute, 'Multideck logo — replaceable Word layout image')
    for shape in root.xpath('//v:shape[@id="issuer"]', namespaces=NS):
        for paragraph in shape.xpath('.//w:txbxContent/w:p', namespaces=NS)[1:]:
            set_text(paragraph, '')
    for paragraph in root.xpath('//w:p', namespaces=NS):
        text = ''.join(paragraph.xpath('./w:r/w:t/text()', namespaces=NS))
        if text in ('AS AGENTS FOR THE CARRIER', 'SHIPPED ON BOARD', 'according to the declaration of the consignor'):
            set_text(paragraph, '')
        elif '{d.fiata.goods[' in text or '{d.fiata.equipment[' in text:
            # Existing reference form loops are a bounded cover-page summary.
            # Full arrays live only in the inline schedule appended below.
            set_text(paragraph, text.replace('[i+1]', '[1]').replace('[i]', '[0]'))


def build(source, target, air=False):
    with zipfile.ZipFile(source) as original:
        root = etree.fromstring(original.read('word/document.xml'))
        patch_air(root) if air else patch_fiata(root)
        body = root.find(qn('w:body'))
        if source.name.startswith('MNG_'):
            # Its public copy label is baked into the form artwork, unlike the
            # editable six-copy MAWB label. Cover that one label in Word without
            # altering the retained bitmap, then put an explicit Draft label in
            # the same slot. Hosted all-page marking remains mandatory as well.
            anchor = body.find(qn('w:p'))
            anchor.append(etree.fromstring(f'''<w:r xmlns:w="{W}" xmlns:v="urn:schemas-microsoft-com:vml"><w:pict>
              <v:shape id="draft-copy-label" type="#_x0000_t202" style="position:absolute;margin-left:155pt;margin-top:773pt;width:300pt;height:24pt;z-index:251659266;mso-position-horizontal-relative:page;mso-position-vertical-relative:page;mso-wrap-style:none" filled="t" fillcolor="#ffffff" stroked="f"><v:fill color="#ffffff"/><v:textbox inset="0,0,0,0"><w:txbxContent><w:p><w:pPr><w:spacing w:before="0" w:after="0"/><w:jc w:val="center"/><w:shd w:val="clear" w:fill="FFFFFF"/></w:pPr><w:r><w:rPr><w:rFonts w:ascii="Arial" w:hAnsi="Arial"/><w:sz w:val="20"/><w:color w:val="FF0000"/></w:rPr><w:t>DRAFT - FOR CUSTOMER REVIEW</w:t></w:r></w:p></w:txbxContent></v:textbox></v:shape>
              </w:pict></w:r>'''))
        original_section = body.find(qn('w:sectPr'))
        assert original_section is not None
        body.remove(original_section)
        section_break = OxmlElement('w:p')
        properties = OxmlElement('w:pPr')
        properties.append(deepcopy(original_section))
        section_break.append(properties)
        body.append(section_break)
        for item in schedule():
            body.append(deepcopy(item))
        with zipfile.ZipFile(target, 'w', zipfile.ZIP_DEFLATED) as output:
            for part in original.infolist():
                output.writestr(part, etree.tostring(root, xml_declaration=True, encoding='UTF-8', standalone=True)
                                if part.filename == 'word/document.xml' else original.read(part.filename))
    with zipfile.ZipFile(source) as original, zipfile.ZipFile(target) as final:
        assert original.namelist() == final.namelist()
        for name in original.namelist():
            if name != 'word/document.xml':
                assert original.read(name) == final.read(name), f'Preserve-only package part changed: {name}'
    print(target.name, hashlib.sha256(target.read_bytes()).hexdigest())


if __name__ == '__main__':
    for name in ('Master_Air_Waybill', 'MNG_Air_Waybill'):
        build(OUT / f'{name}_demo_safe.docx', OUT / f'{name}_booking_draft.docx', air=True)
    build(OUT / 'FIATA_Waybill_demo_safe.docx', OUT / 'FIATA_Waybill_booking_draft.docx')
    build(ROOT / 'supabase/templates/fiata-bill-of-lading/fiata-bill-of-lading-reference-draft.docx',
          OUT / 'FIATA_BOL_reference_booking_draft.docx')
