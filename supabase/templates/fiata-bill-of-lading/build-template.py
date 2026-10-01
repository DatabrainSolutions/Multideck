"""Build an editable, reference-shaped FIATA draft. Never publishes or issues a B/L."""
import argparse
import base64
import copy
import hashlib
import json
import re
from pathlib import Path

import pypdfium2 as pdfium
from docx import Document
from docx.document import Document as DocumentClass
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_ROW_HEIGHT_RULE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt, RGBColor
from docx.text.run import Run
from pypdf import PdfReader

ROOT = Path(__file__).resolve().parent
FORM_WIDTH = 186


def tag(path):
    multiline = path in ['fiata.consignor', 'fiata.consignee', 'fiata.notify', 'fiata.deliveryAgent', 'fiata.remarks'] or path.endswith('.description')
    return "{d." + path + (':convCRLF' if multiline else '') + "}"


def paragraph(cell, text="", size=8, bold=False, align=None):
    p = cell.paragraphs[0] if not cell.paragraphs[0].text and not cell.paragraphs[0].runs else cell.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1
    if align is not None:
        p.alignment = align
    for part in re.split(r'(\{d\.[^{}]+\})', text):
        if not part:
            continue
        chunks = [part[0], part[1:]] if part.startswith('{d.') else [part]
        for i, chunk in enumerate(chunks):
            run = p.add_run(chunk)
            run.font.size = Pt(1 if i else size)
            run.font.color.rgb = RGBColor(0, 0, 0)
            run.bold = bold
    return p


def border(cell, edges=("bottom",), value="single"):
    properties = cell._tc.get_or_add_tcPr()
    borders = properties.find(qn("w:tcBorders"))
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        properties.append(borders)
    for edge in edges:
        node = OxmlElement("w:" + edge)
        node.set(qn("w:val"), value)
        node.set(qn("w:sz"), "6")
        node.set(qn("w:color"), "000000")
        borders.append(node)


def table(parent, widths, rows=1):
    t = parent.add_table(rows=rows, cols=len(widths))
    t.autofit = False
    for column, width in zip(t.columns, widths):
        column.width = Mm(width)
    for row in t.rows:
        props = row._tr.get_or_add_trPr()
        props.append(OxmlElement("w:cantSplit"))
        for c, width in zip(row.cells, widths):
            c.width = Mm(width)
            c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
            margins = OxmlElement("w:tcMar")
            for edge, amount in [("top", 0), ("bottom", 0), ("left", 36), ("right", 36)]:
                n = OxmlElement("w:" + edge)
                n.set(qn("w:w"), str(amount))
                n.set(qn("w:type"), "dxa")
                margins.append(n)
            c._tc.get_or_add_tcPr().append(margins)
            for p in c.paragraphs:
                p.paragraph_format.space_before = Pt(0)
                p.paragraph_format.space_after = Pt(0)
                p.paragraph_format.line_spacing = Pt(1)
                p.add_run().font.size = Pt(1)
    if not isinstance(parent, DocumentClass):
        # Word requires a paragraph after a nested table; keep it unobtrusive.
        parent.paragraphs[-1].paragraph_format.line_spacing = Pt(1)
        parent.paragraphs[-1].paragraph_format.space_after = Pt(0)
        for r in parent.paragraphs[-1].runs:
            r.font.size = Pt(1)
    return t


def height(row, mm):
    row.height = Mm(mm)
    row.height_rule = WD_ROW_HEIGHT_RULE.AT_LEAST


def field(cell, label, path, size=8, label_size=7):
    paragraph(cell, label, label_size, True)
    paragraph(cell, tag(path), size)


def spacer(parent, mm):
    p = parent.add_paragraph()
    p.paragraph_format.line_spacing = Pt(1)
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Mm(mm)
    p.add_run().font.size = Pt(1)


def picture(cell, path, width, height_mm=None, alt=None):
    p = paragraph(cell)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    shape = p.add_run().add_picture(str(path), width=Mm(width), height=Mm(height_mm) if height_mm else None)
    if alt:
        shape._inline.docPr.set("descr", alt)
        shape._inline.docPr.set("title", "Replaceable company logo")


def assets(reference, output):
    output.mkdir(parents=True, exist_ok=True)
    pdf = pdfium.PdfDocument(str(reference))
    rendered = pdf[0].render(scale=3).to_pil()
    # Only the printed form marks are retained. No bird, customer values or signature.
    rendered.crop(tuple(int(v * 3) for v in (310, 9, 371, 38))).save(output / "bifa.png")
    rendered.crop(tuple(int(v * 3) for v in (491, 35, 544, 65))).save(output / "icc.png")
    terms = PdfReader(str(reference)).pages[1].images
    if len(terms) != 1:
        raise ValueError("Expected the supplied conditions page to contain one retained image")
    terms[0].image.save(output / "standard-conditions.png")


def build(reference):
    assets_dir = ROOT / "assets"
    assets(reference, assets_dir)
    d = Document()
    sec = d.sections[0]
    sec.page_width, sec.page_height = Mm(210), Mm(297)
    sec.top_margin, sec.bottom_margin = Mm(4), Mm(7)
    sec.left_margin = sec.right_margin = Mm(12)
    sec.header_distance = sec.footer_distance = Mm(1)
    normal = d.styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(8)
    normal.paragraph_format.space_after = Pt(0)
    normal.paragraph_format.line_spacing = 1
    normal._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "Arial")
    normal._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "Arial")
    d.core_properties.title = "FIATA Bill of Lading"
    d.core_properties.author = "Multideck"

    top = table(d, [94, 92])
    left, right = top.rows[0].cells
    border(left, ("right",))
    parties = table(left, [93], 3)
    for row, label, name in zip(parties.rows, ["Consignor", "Consigned to order of", "Notify Address"], ["consignor", "consignee", "notify"]):
        field(row.cells[0], label, f"fiata.{name}")
        height(row, 21)
        border(row.cells[0])
    routing = table(left, [46.5, 46.5], 3)
    field(routing.cell(0, 1), "Place of Receipt", "fiata.placeOfReceipt", 7)
    field(routing.cell(1, 0), "Ocean Vessel", "fiata.oceanVessel", 7)
    field(routing.cell(1, 1), "Port of loading", "fiata.portOfLoading", 7)
    field(routing.cell(2, 0), "Port of discharge", "fiata.portOfDischarge", 7)
    field(routing.cell(2, 1), "Place of Delivery", "fiata.placeOfDelivery", 7)
    for row in routing.rows:
        height(row, 8)
        for cell in row.cells:
            border(cell)

    heading = table(right, [22, 18, 38, 12])
    picture(heading.cell(0, 0), assets_dir / "bifa.png", 21)
    paragraph(heading.cell(0, 1), "FBL", 13, True)
    paragraph(heading.cell(0, 2), tag("fiata.reference"), 9, align=WD_ALIGN_PARAGRAPH.RIGHT)
    field(heading.cell(0, 3), "Country\nCode", "fiata.countryCode", 8, 5)
    border(heading.cell(0, 2), ("top", "bottom", "left", "right"))
    border(heading.cell(0, 3), ("top", "bottom", "right"))
    title = table(right, [66, 24])
    paragraph(title.cell(0, 0), "NEGOTIABLE FIATA\nMULTIMODAL TRANSPORT\nBILL OF LADING", 9)
    picture(title.cell(0, 1), assets_dir / "icc.png", 20)
    paragraph(right, "issued subject to UNCTAD/ICC Rules for\nMultimodal Transport Documents (ICC Publication 481).", 7)
    spacer(right, 3)
    logo_panel = table(right, [90])
    height(logo_panel.rows[0], 34)
    logo_panel.cell(0, 0).vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    logo = ROOT.parents[2] / "multideck.client/public/email/multideck-logo.png"
    picture(logo_panel.cell(0, 0), logo, 77, 23, "{d.company.logoDataUri:imageFit(contain)} {d.company.logoDataUri:ifEM():drop(img)}")
    paragraph(right, tag("fiata.carrierName"), 16, align=WD_ALIGN_PARAGRAPH.CENTER)
    p = paragraph(right, "BILL OF LADING", 23, True, WD_ALIGN_PARAGRAPH.CENTER)
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.keep_with_next = False
    for r in p.runs:
        r.font.name, r.font.size, r.bold = "Arial", Pt(23), True
    paragraph(right, tag("fiata.billType"), 10, True)
    paragraph(right, tag("fiata.billTypeDescription"), 8)
    for cell in top.rows[0].cells:
        border(cell)
    spacer(d, 1)

    cargo_area = table(d, [FORM_WIDTH])
    height(cargo_area.rows[0], 83)
    body = cargo_area.cell(0, 0)
    goods = table(body, [46, 47, 46, 24, 21], 3)
    labels = ["Marks and Numbers", "Number and Kind of Packages", "Description of Goods", "Gross Weight", "Measurement"]
    paths = ["marks", "packages", "description", "grossWeight", "measurement"]
    for c, label in zip(goods.rows[0].cells, labels):
        paragraph(c, label, 7, True)
    for i in range(2):
        for c, path in zip(goods.rows[i+1].cells, paths):
            paragraph(c, tag(f"fiata.goods[{'i' if i == 0 else 'i+1'}].{path}"), 8)
    spacer(body, 2)
    equipment = table(body, [31, 31, 18, 34, 70], 3)
    for c, label in zip(equipment.rows[0].cells, ["Container", "Seals", "Type", "Packages", "Mode"]):
        paragraph(c, label, 7, True)
    for i in range(2):
        for c, path in zip(equipment.rows[i+1].cells, ["number", "seal", "type", "packages", "mode"]):
            paragraph(c, tag(f"fiata.equipment[{'i' if i == 0 else 'i+1'}].{path}"), 8)
    summary = table(d, [70, 70, 46], 3)
    paragraph(summary.cell(0, 0), tag("fiata.remarks"), 8)
    paragraph(summary.cell(0, 1), "Consol Ref: " + tag("fiata.consolidationReference"), 8)
    paragraph(summary.cell(0, 2), tag("fiata.serviceMode"), 8)
    paragraph(summary.cell(1, 0), "Total Packages: " + tag("fiata.totalPackages"), 8, True)
    paragraph(summary.cell(1, 2), tag("fiata.loadAndCount"), 7, align=WD_ALIGN_PARAGRAPH.RIGHT)
    paragraph(summary.cell(2, 0), tag("fiata.shippedOnBoardStatement"), 7)
    paragraph(summary.cell(2, 1), "according to the declaration of the consignor", 6.5, True)
    for cell in summary.rows[-1].cells:
        border(cell)

    declarations = table(d, [93, 93])
    field(declarations.cell(0, 0), "Declaration of Interest of the consignor in\ntimely delivery (Clause 6.2.)", "fiata.deliveryInterest", 8, 7)
    field(declarations.cell(0, 1), "Declared value for ad valorem rate according to\nthe declaration of the consignor (Clauses 7 and 8)", "fiata.declaredValue", 8, 7)
    height(declarations.rows[0], 14)
    spacer(d, 2)
    for text in [
        "The goods and instructions are accepted and dealt with subject to the Standard Conditions printed overleaf.",
        "Taken in charge in apparent good order and condition, unless otherwise noted herein, at the place of receipt for transport and delivery as mentioned above.",
        "One of these Multimodal Transport Bills of Lading must be surrendered duly endorsed in exchange for the goods. In Witness whereof the original Multimodal Transport Bills of Lading and all of this tenor and date have been signed in the number stated below, one of which being accomplished the other(s) to be void.",
    ]:
        p = d.add_paragraph()
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.line_spacing = 1
        r = p.add_run(text)
        r.font.size, r.bold = Pt(6.5), True

    bottom = table(d, [71, 44, 71], 3)
    field(bottom.cell(0, 0), "Freight amount", "fiata.freightAmount")
    field(bottom.cell(0, 1), "Freight Payable at", "fiata.freightPayableAt", 7)
    field(bottom.cell(0, 2), "Place and date of issue", "fiata.placeAndDateOfIssue", 7)
    field(bottom.cell(1, 0), "Cargo Insurance through the undersigned", "fiata.insurance", 7)
    field(bottom.cell(1, 1), "No. of Originals", "fiata.originals")
    sig = bottom.cell(1, 2).merge(bottom.cell(2, 2))
    paragraph(sig, "Stamp and Signature", 7, True)
    spacer(sig, 16)
    paragraph(sig, tag("fiata.signingCapacity"), 7)
    delivery = bottom.cell(2, 0).merge(bottom.cell(2, 1))
    field(delivery, "For delivery of goods please apply to:", "fiata.deliveryAgent", 8)
    height(bottom.rows[0], 8)
    height(bottom.rows[1], 8)
    height(bottom.rows[2], 23)
    for row in bottom.rows:
        for cell in row.cells:
            border(cell, ("top", "left"))
    p = d.add_paragraph("HBoL-FIATA")
    p.runs[0].font.size = Pt(4)
    d.add_page_break()
    p = d.add_paragraph()
    p.paragraph_format.space_after = Pt(0)
    p.add_run().add_picture(str(assets_dir / "standard-conditions.png"), width=Mm(FORM_WIDTH))
    path = ROOT / "fiata-bill-of-lading-reference-draft.docx"
    d.save(path)
    return path


def fixture():
    logo = ROOT.parents[2] / "multideck.client/public/email/multideck-logo.png"
    return {"company": {"logoDataUri": "data:image/png;base64," + base64.b64encode(logo.read_bytes()).decode()}, "fiata": {
        "reference": "DEMO-FBL-001", "countryCode": "GB",
        "consignor": "EXAMPLE MANUFACTURING LTD\nUNIT 3\nINDUSTRIAL WAY\nREDDITCH\nB98 7XX\nUNITED KINGDOM",
        "consignee": "EXAMPLE TRADING CO LTD\n29 EXAMPLE STREET\nKYOTO\n6060000\nJAPAN\noperations@example.test",
        "notify": "EXAMPLE TRADING CO LTD\n29 EXAMPLE STREET\nKYOTO\n6060000\nJAPAN\noperations@example.test",
        "placeOfReceipt": "REDDITCH, UNITED KINGDOM", "oceanVessel": "EXAMPLE VESSEL / EX001",
        "portOfLoading": "SOUTHAMPTON, UNITED KINGDOM", "portOfDischarge": "NAGOYA, JAPAN", "placeOfDelivery": "NAGOYA, JAPAN",
        "carrierName": "Example Freight Line", "billType": "", "billTypeDescription": "",
        "goods": [{"marks": "EXAMPLE-001", "packages": "1 × 40HC CONTAINER", "description": "COMPLETE PET FOOD\n1,756 PACKAGES\nHS CODE 23091011", "grossWeight": "17944.940 KG", "measurement": "25.000 M3"}],
        "equipment": [{"number": "NOT YET RECORDED", "seal": "NOT YET RECORDED", "type": "40HC", "packages": "20 PLT", "mode": "DOOR/CY"}],
        "remarks": "", "consolidationReference": "DEMO-C001", "serviceMode": "DOOR/CY", "totalPackages": "ONE CONTAINER(S)",
        "loadAndCount": "", "shippedOnBoardStatement": "", "deliveryInterest": "", "declaredValue": "", "freightAmount": "",
        "freightPayableAt": "", "placeAndDateOfIssue": "", "insurance": "", "originals": "",
        "signingCapacity": "", "deliveryAgent": "EXAMPLE DESTINATION AGENT\n1 EXAMPLE STREET\nNAGOYA\nJAPAN",
    }}


def multiple_rows_fixture():
    data = fixture()
    mark = ROOT.parents[2] / 'multideck.client/src/assets/brand/multideck-logo-mark.svg'
    data['company']['logoDataUri'] = 'data:image/svg+xml;base64,' + base64.b64encode(mark.read_bytes()).decode()
    data['fiata']['goods'].extend([
        {'marks': 'EXAMPLE-002', 'packages': '10 CARTONS', 'description': 'DEMO SPARE PARTS', 'grossWeight': '500 KG', 'measurement': '2 M3'},
        {'marks': 'EXAMPLE-003', 'packages': '2 PALLETS', 'description': 'DEMO MACHINE PARTS', 'grossWeight': '120 KG', 'measurement': '1 M3'},
    ])
    data['fiata']['equipment'] = [
        {'number': 'EXAU1234567', 'seal': 'DEMO-SEAL-1', 'type': '40HC', 'packages': '20 PLT', 'mode': 'DOOR/CY'},
        {'number': 'EXAU7654321', 'seal': 'DEMO-SEAL-2', 'type': '20GP', 'packages': '10 CTN', 'mode': 'DOOR/CY'},
    ]
    data['fiata']['totalPackages'] = 'TWO CONTAINER(S)'
    return data


def filled_template(path, data, output, logo_path=None):
    """Local layout QA only. Connected Carbone preview remains the tag-engine check."""
    doc = Document(path)
    root = doc.element
    for row in list(root.iter(qn("w:tr"))):
        text = "".join(n.text or "" for cell in row.findall(qn('w:tc')) for p in cell.findall(qn('w:p')) for n in p.iter(qn('w:t')))
        if '[i+1]' in text:
            row.getparent().remove(row)
        elif '[i]' in text:
            array_path = re.search(r'\{d\.([a-zA-Z.]+)\[i\]', text).group(1)
            items = data
            for segment in array_path.split('.'):
                items = items[segment]
            parent = row.getparent()
            position = parent.index(row)
            for index in range(len(items)):
                clone = copy.deepcopy(row)
                for n in clone.iter(qn('w:t')):
                    n.text = (n.text or '').replace('[i]', f'[{index}]')
                parent.insert(position + index, clone)
            parent.remove(row)
    for p in root.iter(qn('w:p')):
        texts = list(p.iter(qn('w:t')))
        if not texts:
            continue
        text = ''.join(n.text or '' for n in texts)
        if '{d.' not in text:
            continue
        def replace(m):
            value = data
            for segment in re.sub(r'\[(\d+)\]', r'.\1', m.group(1).split(':')[0]).split('.'):
                value = value[int(segment)] if isinstance(value, list) else value.get(segment, "")
            return str(value)
        for n in texts[1:]:
            n.text = ''
        Run(texts[0].getparent(), None).text = re.sub(r"\{d\.([^{}]+)\}", replace, text)
    if logo_path:
        for shape in doc.inline_shapes:
            if "logoDataUri" in shape._inline.docPr.get("descr", ""):
                relationship_id = shape._inline.graphic.graphicData.pic.blipFill.blip.embed
                doc.part.related_parts[relationship_id]._blob = Path(logo_path).read_bytes()
    doc.save(output)


if __name__ == "__main__":
    parser = argparse.ArgumentParser()
    parser.add_argument("reference", type=Path)
    args = parser.parse_args()
    path = build(args.reference)
    data = fixture()
    (ROOT / "sample-data.json").write_text(json.dumps(data, indent=2) + "\n")
    (ROOT / 'sample-multiple-rows.json').write_text(json.dumps(multiple_rows_fixture(), indent=2) + '\n')
    qa = ROOT / "qa"
    qa.mkdir(exist_ok=True)
    filled_template(path, data, qa / "filled.docx")
    print(json.dumps({"template": str(path), "sha256": hashlib.sha256(path.read_bytes()).hexdigest()}))
