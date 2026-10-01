"""Build review-only Word/Carbone layouts and a row-complete BoxTop crosswalk.

The supplied workbook and examples are read-only. Real sample text never enters
the templates or fixtures. No database, Carbone, publish or business-record write
is performed by this builder.
"""
from __future__ import annotations

import argparse
import hashlib
import json
import re
from collections import Counter
from pathlib import Path

import openpyxl
from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt, RGBColor

ROOT = Path(__file__).resolve().parent
WIDTH = 180

# Values are presentation labels supplied by the future authorised source
# adapter, not calculations performed in a document. All units are explicit.
TABLES = {
    "cargo": [("Line", "line", 10), ("Goods and handling", "description", 60), ("Packages", "packages", 27), ("Gross kg", "grossKg", 24), ("CBM", "volume", 20), ("Marks", "marks", 39)],
    "equipment": [("Container or equipment", "reference", 40), ("Type", "type", 22), ("Seal", "seal", 30), ("Packages", "packages", 30), ("Loaded gross kg", "loadedGrossKg", 31), ("VGM kg", "vgmKg", 27)],
    "routing": [("Leg", "line", 10), ("Mode", "mode", 20), ("From", "origin", 38), ("To", "destination", 38), ("Planned departure", "departure", 37), ("Planned arrival", "arrival", 37)],
    "charges": [("Description", "description", 86), ("Quantity", "quantity", 22), ("Unit price", "unitPrice", 36), ("Amount", "amount", 36)],
    "invoiceLines": [("Description", "description", 69), ("Quantity", "quantity", 19), ("Net amount", "net", 31), ("Tax rate", "taxRate", 23), ("Tax", "tax", 19), ("Total", "total", 19)],
    "payments": [("Invoice reference", "invoice", 54), ("Invoice date", "date", 32), ("Original amount", "amount", 34), ("Applied amount", "applied", 30), ("Balance", "balance", 30)],
    "customsCargo": [("Line", "line", 10), ("Goods", "description", 38), ("Packages", "packages", 23), ("Gross kg", "grossKg", 21), ("Net kg", "netKg", 21), ("HS code", "commodity", 25), ("Origin", "originCountry", 14), ("Value", "value", 28)],
    "attachments": [("Document", "name", 72), ("Type", "type", 47), ("Source reference", "reference", 61)],
    "manifest": [("House or job", "reference", 31), ("Shipper", "shipper", 34), ("Consignee", "consignee", 34), ("Goods", "description", 35), ("Packages", "packages", 25), ("Gross kg", "grossKg", 21)],
    "allocations": [("Cargo line", "cargo", 48), ("Container", "equipment", 43), ("Packages", "packages", 27), ("Cargo kg", "grossKg", 25), ("Scope", "scope", 37)],
    "tradeLines": [("Goods", "description", 48), ("Quantity", "quantity", 18), ("Unit", "unit", 19), ("Origin", "originCountry", 17), ("HS code", "commodity", 27), ("Unit value", "unitPrice", 25), ("Value", "amount", 26)],
    "stock": [("SKU", "sku", 25), ("Description", "description", 52), ("Location", "location", 28), ("Quantity", "quantity", 22), ("Unit", "unit", 19), ("Lot or batch", "batch", 34)],
    "vehicles": [("VIN or chassis", "vin", 42), ("Make and model", "model", 40), ("Registration", "registration", 28), ("Mileage", "mileage", 22), ("Condition or exceptions", "condition", 48)],
}

FIXTURE_ROWS = {
    "cargo": [{"line":"1", "description":"Fictional machine components\nFragile", "packages":"40 Pallets", "grossKg":"4,000", "volume":"12.5", "marks":"DEMO ONLY"}],
    "equipment": [{"reference":"Container 1", "type":"40GP", "seal":"Not recorded", "packages":"20 Pallets", "loadedGrossKg":"Not recorded", "vgmKg":"Not recorded"}, {"reference":"Container 2", "type":"40GP", "seal":"Not recorded", "packages":"20 Pallets", "loadedGrossKg":"Not recorded", "vgmKg":"Not recorded"}],
    "routing": [{"line":"1", "mode":"Sea", "origin":"GBFXT", "destination":"INNSA", "departure":"30 Oct 2026", "arrival":"28 Nov 2026"}],
    "charges": [{"description":"Fictional freight service", "quantity":"1", "unitPrice":"GBP 200.00", "amount":"GBP 200.00"}],
    "invoiceLines": [{"description":"Fictional freight service", "quantity":"1", "net":"200.00", "taxRate":"20%", "tax":"40.00", "total":"240.00"}],
    "payments": [{"invoice":"DEMO-INVOICE-001", "date":"30 Sep 2026", "amount":"GBP 240.00", "applied":"GBP 120.00", "balance":"GBP 120.00"}],
    "customsCargo": [{"line":"1", "description":"Fictional components", "packages":"40 Pallets", "grossKg":"4,000", "netKg":"3,800", "commodity":"Not recorded", "originCountry":"GB", "value":"GBP 8,000.00"}],
    "attachments": [{"name":"Fictional invoice.pdf", "type":"Commercial invoice", "reference":"DEMO-TRADE-001"}],
    "manifest": [{"reference":"DEMO-HOUSE-001", "shipper":"Example Exporter", "consignee":"Example Importer", "description":"Machine components", "packages":"40 Pallets", "grossKg":"4,000"}],
    "allocations": [{"cargo":"1 · Machine components", "equipment":"Container 1 · 40GP", "packages":"20 Pallets", "grossKg":"Not recorded", "scope":"Whole journey"}, {"cargo":"1 · Machine components", "equipment":"Container 2 · 40GP", "packages":"20 Pallets", "grossKg":"Not recorded", "scope":"Whole journey"}],
    "tradeLines": [{"description":"Fictional machine components", "quantity":"40", "unit":"Pallets", "originCountry":"GB", "commodity":"Not recorded", "unitPrice":"GBP 200.00", "amount":"GBP 8,000.00"}],
    "stock": [{"sku":"DEMO-SKU-001", "description":"Fictional machine components", "location":"A01-01", "quantity":"40", "unit":"Pallets", "batch":"DEMO-BATCH-001"}],
    "vehicles": [{"vin":"DEMO-CHASSIS-NOT-VALID", "model":"Example vehicle", "registration":"Not recorded", "mileage":"Not recorded", "condition":"Not recorded"}],
}

def key(label: str) -> str:
    words = re.findall(r"[A-Za-z0-9]+", label)
    return words[0].lower() + "".join(w.title() for w in words[1:])

def sample_value(label: str) -> str:
    if label == "Currency": return "GBP"
    if label == "Version": return "2"
    if label in ("Direction", "Mode", "Shipment type", "Status", "Incoterms"):
        return {"Direction":"Export", "Mode":"Sea", "Shipment type":"FCL", "Status":"Planning", "Incoterms":"FOB"}[label]
    if "email" in label.lower(): return "operator@example.test"
    if "name" in label.lower() or label in ("Prepared by", "Checked by", "Picked by", "Owner", "Assigned to", "Driver"):
        return "Example Operator"
    if "reference" in label.lower() or "number" in label.lower(): return "DEMO-001"
    if "date" in label.lower() or label.startswith("Planned"): return "30 Oct 2026"
    return "Not recorded"

def fixture(spec):
    data = {"issuer":{"name":"Example Freight Ltd", "address":"Example office address", "email":"operator@example.test", "phone":"Not recorded"}, "reference":"DEMO-001", "issuedAt":"30 Sep 2026", "fields":{key(f):sample_value(f) for f in spec["fields"]}, "instructions":"Handle with care. Dates shown are planned unless explicitly recorded as actual.", "handling":"Fragile", "terms":"Not recorded", "totals":{"net":"GBP 200.00", "tax":"GBP 40.00", "total":"GBP 240.00"}, "paymentInstructions":"Use the verified Finance payment instructions", "preparedBy":"Example Operator"}
    for party, name in (("customer","Example Customer"), ("shipper","Example Exporter"), ("consignee","Example Importer"), ("notify","Example Agent")):
        data[party] = {"name":name, "address":"Fictional address for template review", "registration":"Not recorded"}
    data["collection"] = {"address":"Explicit fictional collection location", "date":"30 Oct 2026"}
    data["delivery"] = {"address":"Explicit fictional delivery location", "date":"28 Nov 2026"}
    data["route"] = {"origin":"GBFXT", "destination":"INNSA", "mode":"Sea"}
    for table in spec["tables"]: data[table] = FIXTURE_ROWS[table]
    if spec["kind"] == "tradeInvoice":
        # A commercial goods invoice is not the forwarder's service invoice.
        data["issuer"] = {"name":"Example Exporter", "address":"Fictional seller address", "email":"seller@example.test", "phone":"Not recorded"}
        data["customer"] = {"name":"Example Importer", "address":"Fictional buyer address", "registration":"Not recorded"}
        data["totals"] = {"net":"GBP 8,000.00", "tax":"GBP 0.00", "total":"GBP 8,000.00"}
        data["instructions"] = "Working draft for customer approval; not evidence of a customer-issued commercial invoice."
    return {"document":data}

def tag(path): return "{d.document." + path + "}"

def add_value(paragraph, value, label="", template=False, size=9):
    if label:
        run = paragraph.add_run(label.upper() + "\n")
        run.font.size, run.bold = Pt(8), True
    if template and value.startswith("{"):
        # Carbone inherits the first brace's formatting; keep the rest of an
        # editable tag compact without shrinking the generated output.
        paragraph.add_run(value[0]).font.size = Pt(size)
        paragraph.add_run(value[1:]).font.size = Pt(5)
    else:
        paragraph.add_run(value).font.size = Pt(size)

def table(doc, widths, headers=None):
    assert sum(widths) == WIDTH
    t = doc.add_table(rows=1, cols=len(widths)); t.autofit = False
    for col, width in zip(t.columns, widths): col.width = Mm(width)
    for cell, width in zip(t.rows[0].cells, widths): cell.width = Mm(width)
    if headers:
        props = t.rows[0]._tr.get_or_add_trPr()
        repeat = OxmlElement("w:tblHeader"); repeat.set(qn("w:val"),"true"); props.append(repeat)
        for cell, label in zip(t.rows[0].cells, headers):
            add_value(cell.paragraphs[0], label, size=8)
            cell.paragraphs[0].runs[0].bold = True
            shade = OxmlElement("w:shd"); shade.set(qn("w:fill"), "F0F3F2"); cell._tc.get_or_add_tcPr().append(shade)
    return t

def style_table(t):
    for row in t.rows:
        props = row._tr.get_or_add_trPr(); props.append(OxmlElement("w:cantSplit"))
        for cell in row.cells:
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p = cell._tc.get_or_add_tcPr()
            borders = OxmlElement("w:tcBorders")
            for edge in ("top","bottom","left","right"):
                line = OxmlElement("w:"+edge)
                for attr, value in (("val","single"),("sz","4"),("color","D9D9D9")): line.set(qn("w:"+attr),value)
                borders.append(line)
            p.append(borders)
            margins = OxmlElement("w:tcMar")
            for edge, value in (("top","70"),("bottom","70"),("start","100"),("end","100")):
                el = OxmlElement("w:"+edge); el.set(qn("w:w"),value); el.set(qn("w:type"),"dxa"); margins.append(el)
            p.append(margins)
            for para in cell.paragraphs:
                para.paragraph_format.space_after = Pt(0)
                para.paragraph_format.line_spacing = 1.05

def heading(doc, text):
    p=doc.add_paragraph(text, "Heading 2")
    p.paragraph_format.space_before, p.paragraph_format.space_after = Pt(8),Pt(3)

def create_doc(spec, destination, preview=False, stress=False):
    doc=Document(); sec=doc.sections[0]
    sec.page_width,sec.page_height=Mm(210),Mm(297)
    sec.left_margin=sec.right_margin=Mm(15)
    sec.top_margin,sec.bottom_margin=Mm(14),Mm(15)
    sec.footer_distance=Mm(6)
    for style in ("Normal","Title","Heading 1","Heading 2"):
        doc.styles[style].font.name="Arial"
        doc.styles[style].font.color.rgb=RGBColor.from_string("000000")
    # The bundled blank Word document carries a blue Title border. Do not let
    # that unrelated default become part of these operational layouts.
    for border in doc.styles.element.xpath('.//w:pBdr'):
        border.getparent().remove(border)
    doc.styles["Normal"].font.size=Pt(9)
    doc.styles["Normal"].paragraph_format.space_after=Pt(3)
    doc.styles["Normal"].paragraph_format.line_spacing=1.1
    doc.styles["Title"].font.size=Pt(18)
    doc.styles["Heading 2"].font.size=Pt(10)
    doc.core_properties.author="Multideck"
    doc.core_properties.title=spec["name"]
    doc.core_properties.subject="Editable review layout; generation wiring pending"
    data=fixture(spec)["document"]
    def val(path):
        if not preview: return tag(path)
        x=data
        for part in path.split("."): x=x[part]
        return str(x)
    p=doc.add_paragraph(); add_value(p,val("issuer.name"),template=not preview,size=12)
    p=doc.add_paragraph(); add_value(p,val("issuer.address"),template=not preview,size=8)
    doc.add_paragraph(spec["name"],"Title")
    p=doc.add_paragraph()
    add_value(p,val("reference"),"Document reference",not preview)
    if spec["kind"] == "invoice" and spec["code"] == "FREIGHT_PROFORMA":
        doc.add_paragraph("Proforma only not a tax invoice")
    meta=table(doc,[60,60,60])
    for i,label in enumerate(spec["fields"]):
        row=i//3
        if row>=len(meta.rows): meta.add_row()
        add_value(meta.cell(row,i%3).paragraphs[0],val("fields."+key(label)),label,not preview)
    style_table(meta)
    financial=spec["kind"] in ("invoice","quote","remittance","tradeInvoice")
    heading(doc,"Bill to" if financial else "Parties")
    parties=table(doc,[90,90])
    party_names=[("customer","Customer"),("issuer","From")] if financial else [("shipper","Shipper"),("consignee","Consignee")]
    for cell,(party,label) in zip(parties.rows[0].cells,party_names):
        add_value(cell.paragraphs[0],val(party+".name"),label,not preview)
        p=cell.add_paragraph();add_value(p,val(party+".address"),template=not preview)
    style_table(parties)
    if spec["kind"] == "quote":
        heading(doc,"Shipment parties")
        parties=table(doc,[90,90])
        for cell,(party,label) in zip(parties.rows[0].cells,[("shipper","Shipper"),("consignee","Consignee")]):
            add_value(cell.paragraphs[0],val(party+".name"),label,not preview)
            add_value(cell.add_paragraph(),val(party+".address"),template=not preview)
        style_table(parties)
    if (not financial and spec["kind"] not in ("warehouse",)) or spec["kind"]=="quote":
        heading(doc,"Movement")
        movement=table(doc,[90,90])
        for cell,(party,label) in zip(movement.rows[0].cells,[("collection","Collection point"),("delivery","Delivery point")]):
            add_value(cell.paragraphs[0],val(party+".address"),label,not preview)
            p=cell.add_paragraph();add_value(p,val(party+".date"),"Planned date",not preview)
        style_table(movement)
    for name in spec["tables"]:
        heading(doc,{"invoiceLines":"Invoice lines", "customsCargo":"Cargo for Customs", "tradeLines":"Trade goods", "allocations":"Cargo load plan", "stock":"Stock lines"}.get(name,name.title()))
        columns=TABLES[name]
        t=table(doc,[c[2] for c in columns],[c[0] for c in columns])
        items=data[name] if preview else [None]
        if stress and preview:
            items=[dict(data[name][0]) for _ in range(20)]
            for i,item in enumerate(items):
                if "line" in item: item["line"]=str(i+1)
                if "description" in item: item["description"]="Long fictional description to check wrapping and repeated table headers without clipping"
        for item in items:
            cells=t.add_row().cells
            for cell,col in zip(cells,columns):
                add_value(cell.paragraphs[0],str(item[col[1]]) if preview else tag(name+"[i]."+col[1]),template=not preview,size=8.5)
        if not preview:
            cells=t.add_row().cells
            add_value(cells[0].paragraphs[0],tag(name+"[i+1]"),template=True,size=8.5)
        style_table(t)
    if spec["kind"] in ("invoice","quote","tradeInvoice"):
        heading(doc,"Totals")
        total=table(doc,[60,60,60])
        for cell,label in zip(total.rows[0].cells,["net","tax","total"]): add_value(cell.paragraphs[0],val("totals."+label),label.title(),not preview)
        style_table(total)
        if spec["kind"]=="invoice":
            p=doc.add_paragraph();add_value(p,val("paymentInstructions"),"Payment instructions",not preview)
    heading(doc,"Notes" if financial and spec["kind"]!="quote" else "Handling and instructions")
    if not financial or spec["kind"]=="quote":
        p=doc.add_paragraph();add_value(p,val("handling"),"Handling",not preview)
    p=doc.add_paragraph();add_value(p,val("instructions"),template=not preview)
    if financial:
        p=doc.add_paragraph();add_value(p,val("terms"),"Terms",not preview)
    footer=sec.footer.paragraphs[0]
    footer.alignment=WD_ALIGN_PARAGRAPH.RIGHT
    footer.add_run("Prepared by ").font.size=Pt(8)
    add_value(footer,val("preparedBy"),template=not preview,size=8)
    doc.save(destination)

def crosswalk(specs, workbook, example_root):
    s=openpyxl.load_workbook(workbook,read_only=True,data_only=True)["Catalogue"]
    row_specs={r:spec for spec in specs for r in spec["rows"]}
    reuse={13:"JOB_CONFIRMATION",24:"JOB_CONFIRMATION",34:"JOB_CONFIRMATION",178:"JOB_CONFIRMATION",22:"FIATA_BOL",29:"MASTER_AIR_WAYBILL",30:"MASTER_AIR_WAYBILL",31:"MASTER_AIR_WAYBILL"}
    controlled={18,19,20,21,26,27,28,32,33,55,56,58,59,60,61,73,179}
    results=[]
    for n,row in enumerate(s.values,1):
        if n<=5: continue
        item={"row":n,"section":row[0],"menu":row[1],"name":row[2],"outputCode":row[3],"evidence":row[4],"mapping":row[5],"notes":row[9]}
        if n in row_specs:
            spec=row_specs[n]
            item.update(status="draft layout created",templateCode="CATALOGUE_"+spec["code"],sourceRecord=spec["source"],gate=spec["gate"],samples=spec["samples"])
        elif n in reuse:
            item.update(status="reuse existing template",templateCode=reuse[n],gate="Existing layout retained. Selection, authorised field projection and specialised issuance remain to be verified; no new publication.")
        elif n in controlled:
            item.update(status="controlled form review required",gate="Needs authorised form/terms, qualified review or verified evidence. Keep supplied originals attached; do not reproduce certification or signatures.")
        else:
            item.update(status="report definition required",gate="Define the Multideck source, filters, aggregation and permission scope before authoring a report. BoxTop output names alone do not specify those rules.")
        results.append(item)
    examples=[p for p in example_root.rglob("*") if p.is_file() and p.suffix.lower() in (".pdf",".docx",".dotx",".xlsx",".tif",".tiff",".txt") and "node_modules" not in p.parts and not p.name.startswith("~$") and ("incoming-documents" in p.parts or "sample-documents" in p.parts)]
    for spec in specs:
        for sample in spec["samples"]:
            if not (example_root/sample).is_file(): raise FileNotFoundError(sample)
    return results,examples

def main():
    parser=argparse.ArgumentParser()
    parser.add_argument("--workbook",default="/Users/leewright/Documents/BoxTop-document-catalogue.xlsx")
    parser.add_argument("--examples",default="/Users/leewright/repo/MultiDeck_Project/multideck-layer0-mcp")
    parser.add_argument("--out",type=Path,default=ROOT)
    parser.add_argument("--refresh-fixture",help="Refresh one fictional fixture without rewriting uploaded Word sources or the manifest")
    args=parser.parse_args();specs=json.loads((ROOT/"catalogue-specs.json").read_text())
    out=args.out;out.mkdir(parents=True,exist_ok=True)
    for sub in ("templates","samples","qa-layouts"): (out/sub).mkdir(exist_ok=True)
    if args.refresh_fixture:
        spec=next(s for s in specs if s["code"]==args.refresh_fixture)
        code="CATALOGUE_"+spec["code"]
        create_doc(spec,out/"qa-layouts"/(code+".docx"),preview=True)
        (out/"samples"/(code+".json")).write_text(json.dumps(fixture(spec),indent=2)+"\n")
        print(json.dumps({"fixtureRefreshed":code,"sourceUnchanged":True}))
        return
    templates=[]
    for spec in specs:
        code="CATALOGUE_"+spec["code"]
        target=out/"templates"/(code+".docx")
        create_doc(spec,target)
        create_doc(spec,out/"qa-layouts"/(code+".docx"),preview=True)
        data=fixture(spec)
        (out/"samples"/(code+".json")).write_text(json.dumps(data,indent=2)+"\n")
        templates.append({"code":code,"name":spec["name"],"status":"draft","sourceRecord":spec["source"],"file":"templates/"+target.name,"sha256":hashlib.sha256(target.read_bytes()).hexdigest(),"catalogueRows":spec["rows"],"samples":spec["samples"],"wiringStatus":"not wired", "reviewGate":spec["gate"]})
    # A multi-page layout fixture checks row growth without deriving unknown weights.
    create_doc(next(s for s in specs if s["code"]=="LOADING_MANIFEST"),out/"qa-layouts"/"STRESS_LOADING_MANIFEST.docx",preview=True,stress=True)
    rows,examples=crosswalk(specs,args.workbook,Path(args.examples))
    manifest={"schemaVersion":1,"workbook":"BoxTop-document-catalogue.xlsx","sheet":"Catalogue","range":"A6:J194","catalogueEntries":len(rows),"distinctOutputCodes":len({r["outputCode"] for r in rows}),"referenceFiles":len(examples),"templates":templates,"crosswalk":rows,"summary":dict(Counter(r["status"] for r in rows)),"candidateDataContract":"document namespace; authorised source adapters deliberately deferred"}
    (out/"catalogue-crosswalk.json").write_text(json.dumps(manifest,indent=2,default=str)+"\n")
    print(json.dumps({"templates":len(templates),"catalogueEntries":len(rows),"referenceFiles":len(examples),"summary":manifest["summary"]}))

if __name__=="__main__": main()
